"""Build ml/manifest.csv with perceptual-hash near-duplicate grouping and 70/15/15 splits.

Never puts Uganda or RoCoLe images in train or val.
Usage (from app/ml):  python manifest.py
"""
from __future__ import annotations

import csv
import random
import sys
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

from PIL import Image
import imagehash
from tqdm import tqdm

from common import CAP_AFTER_DEDUPE, HELDOUT_SOURCES, LABELS, ML, RAW, ROOT, SEED

CACHE = RAW / "_cache" / "phash.tsv"
MANIFEST = ML / "manifest.csv"
DATA_MANIFEST = ML / "data_manifest.csv"
HAMMING = 6
EXTS = {".jpg", ".jpeg", ".png", ".JPG", ".JPEG", ".PNG"}

# BRACOL predominant_stress: confirmed against binary columns in dataset.csv
# 0 healthy, 1 miner, 2 rust, 3 phoma, 4 cercospora. No "brown leaf spot" column.
BRACOL_STRESS = {0: "healthy", 1: "miner", 2: "rust", 3: "phoma", 4: "cercospora"}


def _rel(p: Path) -> str:
    try:
        return p.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(p)


def _iter_images(folder: Path):
    if not folder.exists():
        return
    for p in folder.rglob("*"):
        if p.is_file() and p.suffix in EXTS and p.name not in {".extracted", ".exported"}:
            yield p


def collect_rows() -> list[dict]:
    rows: list[dict] = []

    jmuben = [
        (RAW / "jmuben" / "rust", "jmuben", "Leaf rust", "rust"),
        (RAW / "jmuben" / "cercospora", "jmuben", "Cerscospora", "cercospora"),
        (RAW / "jmuben" / "phoma", "jmuben", "Phoma", "phoma"),
        (RAW / "jmuben2" / "healthy", "jmuben2", "Healthy", "healthy"),
        (RAW / "jmuben2" / "miner", "jmuben2", "Miner", "miner"),
    ]
    for folder, src, orig, lab in jmuben:
        for p in _iter_images(folder):
            rows.append(_row(p, src, orig, lab, False))

    rows.extend(_bracol_rows())

    pd = RAW / "plantdoc_images"
    if pd.exists():
        for p in _iter_images(pd):
            orig = p.name.split("__")[1] if "__" in p.name else "plantdoc"
            rows.append(_row(p, "plantdoc", orig, "not_leaf", False))

    syn = RAW / "synthetic_bg"
    if syn.exists():
        for p in _iter_images(syn):
            rows.append(_row(p, "synthetic_bg", p.stem.rsplit("_", 1)[0], "not_leaf", True))

    own = RAW / "own_negatives"
    if own.exists():
        for p in _iter_images(own):
            rows.append(_row(p, "own_negatives", "own", "not_leaf", False))

    # Held-out: never train. Prefix mapping confirmed on 1_* (healthy coffee leaf).
    # 1200_ rust and 2300_ phoma are inferred from the record counts; confirm by eye
    # once those prefixes finish downloading.
    ug = RAW / "uganda"
    if ug.exists():
        for p in _iter_images(ug):
            pref = p.stem.split("_")[0]
            lab = {"1": "healthy", "1200": "rust", "2300": "phoma"}.get(pref)
            if lab is None:
                continue
            rows.append(_row(p, "uganda", f"prefix_{pref}", lab, False))

    ro = RAW / "rocole"
    xlsx = next(ro.glob("*.xlsx"), None) if ro.exists() else None
    labels_by_name: dict[str, tuple[str, str]] = {}
    if xlsx is not None:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
            ws = wb.active
            headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            for rec in ws.iter_rows(min_row=2, values_only=True):
                d = {headers[i]: rec[i] if i < len(rec) else None for i in range(len(headers))}
                name = None
                for k in ("file", "filename", "image", "img"):
                    if d.get(k):
                        name = Path(str(d[k])).name
                        break
                if not name and rec and rec[0]:
                    name = Path(str(rec[0])).name
                if not name:
                    continue
                blob = " ".join(str(v).lower() for v in rec if v is not None)
                if "mite" in blob or "spider" in blob:
                    labels_by_name[name] = ("red_spider_mite", "not_leaf")
                elif "rust" in blob or "unhealthy" in blob:
                    labels_by_name[name] = ("rust", "rust")
                elif "healthy" in blob:
                    labels_by_name[name] = ("healthy", "healthy")
            wb.close()
        except Exception as e:
            print(f"RoCoLe xlsx read failed: {e}", flush=True)
    if ro.exists():
        for p in _iter_images(ro):
            orig, lab = labels_by_name.get(p.name, ("unknown", "unknown"))
            # File-name fallback: C#P#E# often diseased, C#P#H# healthy. Still held-out.
            if lab == "unknown":
                stem = p.stem.upper()
                if "HE" in stem:
                    orig, lab = ("unlabelled", "unknown")
                elif stem[-2:-1] == "H" or stem.endswith("H1") or "H" in stem[-3:]:
                    orig, lab = ("healthy_guess", "healthy")
                else:
                    orig, lab = ("unhealthy_guess", "rust")
            rows.append(_row(p, "rocole", orig, lab, False))

    return rows


def _row(path: Path, source: str, original: str, label: str, synthetic: bool) -> dict:
    return {
        "path": _rel(path),
        "abs": str(path),
        "source": source,
        "original_label": original,
        "label": label,
        "synthetic": "1" if synthetic else "0",
    }


def _bracol_rows() -> list[dict]:
    csv_path = None
    for cand in (RAW / "bracol").rglob("dataset.csv"):
        csv_path = cand
        break
    if csv_path is None:
        print("BRACOL dataset.csv not found; skipping BRACOL", flush=True)
        return []

    images: dict[int, Path] = {}
    for p in _iter_images(RAW / "bracol"):
        if p.stem.isdigit():
            images[int(p.stem)] = p

    out = []
    with csv_path.open(newline="", encoding="utf-8") as f:
        for rec in csv.DictReader(f):
            try:
                i = int(rec["id"])
                stress = int(rec["predominant_stress"])
            except (KeyError, ValueError):
                continue
            lab = BRACOL_STRESS.get(stress)
            if lab is None:
                continue
            p = images.get(i)
            if p is None:
                # Common BRACOL names: 1.jpg or 0001.jpg
                for key in (i,):
                    pass
                continue
            flags = [k for k in ("miner", "rust", "phoma", "cercospora") if rec.get(k) not in (None, "", "0", 0)]
            orig = ",".join(flags) if flags else "healthy"
            out.append(_row(p, "bracol", orig, lab, False))
    print(f"BRACOL labelled images matched: {len(out)} (csv at {csv_path})", flush=True)
    return out


def _hash_one(abs_path: str) -> tuple[str, str | None]:
    try:
        with Image.open(abs_path) as im:
            im = im.convert("RGB")
            return abs_path, str(imagehash.phash(im))
    except Exception:
        return abs_path, None


def load_cache() -> dict[str, str]:
    cache = {}
    if CACHE.exists():
        for line in CACHE.read_text(encoding="utf-8").splitlines():
            if "\t" in line:
                k, v = line.split("\t", 1)
                cache[k] = v
    return cache


def save_cache(cache: dict[str, str]) -> None:
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    tmp = CACHE.with_suffix(".tmp")
    tmp.write_text("".join(f"{k}\t{v}\n" for k, v in cache.items()), encoding="utf-8")
    tmp.replace(CACHE)


def hash_rows(rows: list[dict]) -> None:
    cache = load_cache()
    need = [r for r in rows if r["abs"] not in cache]
    print(f"hash cache {len(cache)} already, {len(need)} to compute", flush=True)
    if need:
        workers = min(8, max(1, (len(need) // 200) or 1))
        done = 0
        with ProcessPoolExecutor(max_workers=workers) as ex:
            futs = {ex.submit(_hash_one, r["abs"]): r for r in need}
            for fut in tqdm(as_completed(futs), total=len(futs), desc="phash"):
                abs_path, h = fut.result()
                if h:
                    cache[abs_path] = h
                done += 1
                if done % 2000 == 0:
                    save_cache(cache)
        save_cache(cache)
    for r in rows:
        r["phash"] = cache.get(r["abs"], "")


class UF:
    def __init__(self, n: int):
        self.p = list(range(n))
        self.r = [0] * n

    def find(self, x: int) -> int:
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a: int, b: int) -> None:
        ra, rb = self.find(a), self.find(b)
        if ra == rb:
            return
        if self.r[ra] < self.r[rb]:
            self.p[ra] = rb
        elif self.r[ra] > self.r[rb]:
            self.p[rb] = ra
        else:
            self.p[rb] = ra
            self.r[ra] += 1


def cluster(rows: list[dict]) -> int:
    """Union hashes with Hamming distance <= 6, within (source, label). Returns pair merges."""
    groups: dict[tuple[str, str], list[int]] = defaultdict(list)
    for i, r in enumerate(rows):
        if r.get("phash"):
            groups[(r["source"], r["label"])].append(i)
    uf = UF(len(rows))
    merges = 0
    for idxs in groups.values():
        ints = []
        for i in idxs:
            try:
                ints.append((i, int(rows[i]["phash"], 16)))
            except ValueError:
                continue
        buckets: dict[tuple[int, int], list[tuple[int, int]]] = defaultdict(list)
        for i, h in ints:
            for band in range(8):
                buckets[(band, (h >> (band * 8)) & 0xFF)].append((i, h))
        seen = set()
        for bucket in buckets.values():
            if len(bucket) < 2 or len(bucket) > 400:
                continue
            for a in range(len(bucket)):
                ia, ha = bucket[a]
                for b in range(a + 1, len(bucket)):
                    ib, hb = bucket[b]
                    pair = (ia, ib) if ia < ib else (ib, ia)
                    if pair in seen:
                        continue
                    seen.add(pair)
                    if (ha ^ hb).bit_count() <= HAMMING:
                        uf.union(ia, ib)
                        merges += 1
    for i, r in enumerate(rows):
        r["cluster_id"] = f"{r['source']}:{r['label']}:{uf.find(i)}"
    return merges


def assign_splits(rows: list[dict]) -> None:
    rng = random.Random(SEED)
    # Held-out stay held-out even if a hash collides with train.
    for r in rows:
        if r["source"] == "uganda":
            r["split"] = "heldout_uganda"
        elif r["source"] == "rocole":
            r["split"] = "heldout_rocole"
        elif r["label"] not in LABELS:
            r["split"] = "unused"

    pool = [r for r in rows if r["source"] not in HELDOUT_SOURCES and r["label"] in LABELS]
    by_cluster: dict[str, list[dict]] = defaultdict(list)
    for r in pool:
        by_cluster[r["cluster_id"]].append(r)

    # Cap oversized classes by sampling whole clusters (seed 42).
    clusters_by_label: dict[str, list[str]] = defaultdict(list)
    for cid, members in by_cluster.items():
        clusters_by_label[members[0]["label"]].append(cid)

    keep_cids = set()
    for lab, cids in sorted(clusters_by_label.items()):
        cids = sorted(cids)
        rng.shuffle(cids)
        cap = CAP_AFTER_DEDUPE.get(lab)
        if cap is None:
            keep_cids.update(cids)
            continue
        n = 0
        for cid in cids:
            if n >= cap:
                for r in by_cluster[cid]:
                    r["split"] = "unused_capped"
                continue
            keep_cids.add(cid)
            n += len(by_cluster[cid])

    # Stratified 70/15/15 on clusters, by class.
    by_lab: dict[str, list[str]] = defaultdict(list)
    for cid in keep_cids:
        by_lab[by_cluster[cid][0]["label"]].append(cid)
    for lab, cids in sorted(by_lab.items()):
        cids = sorted(cids)
        rng.shuffle(cids)
        n = len(cids)
        n_train = int(n * 0.70)
        n_val = int(n * 0.15)
        for i, cid in enumerate(cids):
            if i < n_train:
                sp = "train"
            elif i < n_train + n_val:
                sp = "val"
            else:
                sp = "test"
            for r in by_cluster[cid]:
                r["split"] = sp


def write_manifest(rows: list[dict]) -> None:
    fields = ["path", "source", "original_label", "label", "split", "phash", "cluster_id", "synthetic"]
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    with MANIFEST.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, "") for k in fields})


def write_data_manifest(rows: list[dict]) -> None:
    licences = {
        "jmuben": ("CC BY 4.0", "train/val/test", "Kenya Arabica; rust, cercospora, phoma; cropped and augmented"),
        "jmuben2": ("CC BY 4.0", "train/val/test", "Kenya Arabica; healthy and miner; augmented; capped at 9000 after dedupe"),
        "bracol": ("CC BY 4.0", "train/val/test", "Brazil Arabica; labels from dataset.csv predominant_stress; zip was truncated"),
        "plantdoc": ("CC BY 4.0", "train/val/test", "non-coffee leaves used only as not_leaf"),
        "synthetic_bg": ("synthetic", "train/val/test", "generated paper/wood/soil/skin/grey backgrounds; labelled synthetic"),
        "own_negatives": ("own", "train/val/test", "real non-leaf photos if present"),
        "uganda": ("CC BY 4.0", "held-out test only", "never trained on; augmented; prefix-to-class inferred"),
        "rocole": ("CC BY 4.0", "held-out test only", "never trained on; Robusta field photos; split by plant later"),
    }
    with DATA_MANIFEST.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["dataset", "licence", "role", "n_raw", "n_train", "n_val", "n_test", "n_heldout", "n_unused", "notes"],
        )
        w.writeheader()
        sources = sorted({r["source"] for r in rows})
        for src in sources:
            sub = [r for r in rows if r["source"] == src]
            lic, role, notes = licences.get(src, ("unknown", "unknown", ""))
            w.writerow({
                "dataset": src,
                "licence": lic,
                "role": role,
                "n_raw": len(sub),
                "n_train": sum(1 for r in sub if r.get("split") == "train"),
                "n_val": sum(1 for r in sub if r.get("split") == "val"),
                "n_test": sum(1 for r in sub if r.get("split") == "test"),
                "n_heldout": sum(1 for r in sub if str(r.get("split", "")).startswith("heldout")),
                "n_unused": sum(1 for r in sub if str(r.get("split", "")).startswith("unused")),
                "notes": notes,
            })


def main() -> int:
    print("collecting files", flush=True)
    rows = collect_rows()
    print(f"collected {len(rows)} files  {Counter(r['source'] for r in rows)}", flush=True)
    print(f"labels    {Counter(r['label'] for r in rows)}", flush=True)
    hash_rows(rows)
    hashed = sum(1 for r in rows if r.get("phash"))
    print(f"hashed {hashed}/{len(rows)}", flush=True)
    merges = cluster(rows)
    n_clusters = len({r.get("cluster_id") for r in rows})
    print(f"near-dup merges (pairs within hamming {HAMMING}): {merges}; clusters: {n_clusters}", flush=True)
    assign_splits(rows)
    write_manifest(rows)
    write_data_manifest(rows)
    print("splits", Counter(r.get("split") for r in rows), flush=True)
    print("train labels", Counter(r["label"] for r in rows if r.get("split") == "train"), flush=True)
    print(f"wrote {MANIFEST} and {DATA_MANIFEST}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
